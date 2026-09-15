"""Unit tests for posted-money schedule reconciliation and Pay Bill layout helpers."""

from __future__ import annotations

from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.exceptions import ConflictError
from app.reports.canonical_pay_bill_common import pay_bill_page_capacity_errors
from app.reports.canonical_pay_bill_excel import _footer_text, _money_cell
from app.reports.posted_run import ZERO, reconcile_schedule_totals


def test_reconcile_schedule_totals_accepts_exact_match() -> None:
    reconcile_schedule_totals(
        report_type="hba_schedule",
        posted={"HBA_INSTALLMENT": Decimal("72723.00")},
        emitted={"HBA_INSTALLMENT": Decimal("72723.00")},
    )


def test_reconcile_schedule_totals_accepts_legitimate_exclusions() -> None:
    reconcile_schedule_totals(
        report_type="gpf_mumbai_schedule",
        posted={"GPF_SUBSCRIPTION": Decimal("280000.00")},
        emitted={"GPF_SUBSCRIPTION": Decimal("165000.00")},
        excluded={"GPF_SUBSCRIPTION": Decimal("115000.00")},
    )


def test_reconcile_schedule_totals_normalizes_amounts() -> None:
    reconcile_schedule_totals(
        report_type="gis_schedule",
        posted={"GIS": "22440.0"},
        emitted={"GIS": 22440},
    )


def test_reconcile_schedule_totals_raises_on_dropped_posted_money() -> None:
    with pytest.raises(ConflictError, match="do not reconcile") as exc_info:
        reconcile_schedule_totals(
            report_type="accommodation_mumbai_schedule",
            posted={"ACCOMMODATION_LICENSE_FEE": Decimal("10419.00")},
            emitted={"ACCOMMODATION_LICENSE_FEE": Decimal("9000.00")},
        )
    details = exc_info.value.details or {}
    assert details["error_code"] == "report_schedule_reconciliation"
    assert details["mismatches"]["ACCOMMODATION_LICENSE_FEE"] == {
        "posted": "10419.00",
        "emitted": "9000.00",
        "excluded": "0.00",
    }


def test_reconcile_schedule_totals_flags_unexpected_emitted_money() -> None:
    with pytest.raises(ConflictError, match="do not reconcile"):
        reconcile_schedule_totals(
            report_type="nps_contribution_schedule",
            posted={"NPS_EMPLOYEE": ZERO},
            emitted={"NPS_EMPLOYEE": Decimal("5.00")},
        )


def test_pay_bill_page_capacity_allows_reference_topology() -> None:
    # Reference roster: group boundaries at serials 1, 2, 3, 16, 25, 26.
    groups: list[tuple[str, str, str, str, str]] = []
    for serial in range(1, 29):
        group_index = (
            0
            if serial == 1
            else 1
            if serial == 2
            else 2
            if serial < 16
            else 3
            if serial < 25
            else 4
            if serial == 25
            else 5
        )
        groups.append((f"g{group_index}", f"Post {group_index}", "10", "2", "S-10"))
    assert pay_bill_page_capacity_errors(groups) == []


def test_pay_bill_page_capacity_flags_generic_overflow() -> None:
    # Eight distinct post groups under employees 1-8 overflow the first fixed
    # page: 8 * 6 employee rows + 8 group headers cannot fit before row 62.
    groups = [(f"post-{index}", f"Post {index}", "", "", "") for index in range(8)]
    errors = pay_bill_page_capacity_errors(groups)
    assert len(errors) == 1
    assert "row 67" in errors[0]
    assert "employees 1-8" in errors[0]


def test_pay_bill_page_capacity_flags_second_page_overflow() -> None:
    groups = [(f"post-{index}", f"Post {index}", "", "", "") for index in range(18)]
    errors = pay_bill_page_capacity_errors(groups)
    assert len(errors) == 2
    assert "row 67" in errors[0]
    assert "row 134" in errors[1]
    assert "employees 9-18" in errors[1]


def test_footer_text_escapes_ampersand_and_strips_control_chars() -> None:
    assert _footer_text("R&B Office") == "R&&B Office"
    assert _footer_text("A\x07B") == "AB"
    assert _footer_text("Plain") == "Plain"


def test_money_cell_writes_decimal_not_float() -> None:
    ws = Workbook().active
    assert ws is not None
    _money_cell(ws, 1, 1, Decimal("123.45"))
    assert isinstance(ws.cell(1, 1).value, Decimal)
    assert ws.cell(1, 1).value == Decimal("123.45")
    _money_cell(ws, 2, 1, None)
    assert ws.cell(2, 1).value is None
